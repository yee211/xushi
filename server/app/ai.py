"""通过 OpenAI 兼容协议调用大模型，把 Excel 课表结构化为固定 schema。

配置优先级：管理后台 llm_configs（scope=schedule_import，密钥加密入库）>
.env 的 AI_BASE_URL / AI_API_KEY / AI_MODEL。换厂商只需改 .env 三项。
AI_API_KEY 为空时本模块直接返回 ai-not-configured，调用方会自动回退到
app/parser.py 的确定性解析，因此未配置任何密钥时依然完整可用。
"""
import json
import logging
import os
import re
from pathlib import Path

import httpx
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, ValidationError, field_validator, model_validator

from .excel import read_sheets
from .parser import course_color, merge_section_courses, parse_weeks
from .services.llm_config import get_llm_config

logger = logging.getLogger("classschedule")

AI_TIMEOUT = min(90.0, max(5.0, float(os.getenv("AI_TIMEOUT_SECONDS", "75"))))
AI_MAX_INPUT_CHARS = int(os.getenv("AI_MAX_INPUT_CHARS", "60000"))

# 单个 sheet 的行数上限，防止带大量尾部格式的空表拖垮序列化
MAX_ROWS_PER_SHEET = 1500

SYSTEM_PROMPT = """你是教务课表结构化解析器。用户会提供一个 Excel 工作簿的单元格矩阵文本，格式为逐 sheet 的
`=== SHEET n: "表名" (rows=.., cols=..) ===`，其下可能有 `MERGED:` 行列出合并单元格范围，
再往下每行形如 `ROW <行号>: <坐标>="<值>" <坐标>="<值>" ...`，值内的 `\\n` 表示原单元格中的换行。

你的任务是从中提取课程安排，并且只输出一个 JSON 对象，禁止输出 markdown 代码块、注释或任何解释文字。

输出 JSON 结构：
{
  "name": "课表名称，如“张三的课表”；无法判断时填“我的课表”",
  "term": "学期标识，优先取标题中的“YYYY-YYYY学年第N学期”；确实无法判断时填“导入课表”",
  "courses": [
    {
      "name": "课程名称，必填",
      "teacher": "授课教师，缺失填空字符串",
      "room": "教室或上课地点，缺失填空字符串",
      "weekday": 1,
      "start_section": 1,
      "end_section": 2,
      "weeks": "2-16(双)"
    }
  ]
}

解析规则：
1. weekday：1=周一、2=周二 …… 7=周日。“星期一/周一/礼拜一”等写法一律映射为对应数字。
2. start_section / end_section：1 到 12 的整数，取自“节次”列或行标签（如“第1-2节”“1-2”“上午第1节”）。
3. 连堂必须合并成一条记录：同一课程、同一教师、同一教室且节次连续（如 1-2 节与 3-4 节）时，
   输出 start_section=1、end_section=4，不要拆成两条。
4. weeks：**原样抄录表中的周次表达式字符串，不要展开成数字数组**。例如 "1-16"、"2-16(双)"、
   "3-17(单)"、"1,3,5-8"、"11-15"、"1"。保留单/双标记与括号，可去掉“第”“周”等冗余字。
   若确实无法判断周次，输出空字符串 ""。
5. 同一个单元格里含多门课程时（如同一时段上下两段不同课程），拆成多条记录。
6. 教室必须保留完整全称（包含楼栋编号与前缀，例如 7-南312、4-南407，严禁擅自省略楼栋号）；teacher / room 缺失时输出空字符串，严禁编造，也不要从相邻行推断；单元格为空则不要输出。
7. 无法确定 weekday 或节次的行直接丢弃，不要猜测。
8. 忽略表头行、序号列、说明、备注、合计、统计、签名、审核等非课程信息。
9. 课程名称要清理多余空白与换行，但不要把课程名与教师名粘连在一起。
10. 若整份文件里确实没有任何课程，courses 输出空数组。
11. 输出确定性去重后的结果，同一 (课程,星期,节次,周次) 只能出现一次。"""


class AICourse(BaseModel):
    """与 courses 表字段严格对齐；color 由后端确定性生成，不向模型索要。"""
    name: str
    teacher: str = ""
    room: str = ""
    weekday: int = Field(ge=1, le=7)
    start_section: int = Field(ge=1, le=12)
    end_section: int = Field(ge=1, le=12)
    weeks: list[int] = Field(default_factory=list)

    @field_validator("name", "teacher", "room")
    @classmethod
    def clean_text(cls, value, info):
        # 截断到数据库列宽而不是直接报错，避免个别超长字段拖垮整批导入
        limit = {"name": 80, "teacher": 40, "room": 40}[info.field_name]
        return re.sub(r"\s+", " ", str(value or "")).strip()[:limit]

    @field_validator("weeks", mode="before")
    @classmethod
    def coerce_weeks(cls, value):
        # 提示词要求模型直接回传 "1-16" / "2-16(双)" 这样的原式（输出 token 比展开数组少一个量级，
        # 直接决定了接口耗时），用本地解析器展开；兼容模型仍然返数组的情况
        if isinstance(value, str):
            return parse_weeks(value)
        return value

    @field_validator("weeks")
    @classmethod
    def clean_weeks(cls, value):
        return sorted({week for week in value if isinstance(week, int) and 1 <= week <= 30})

    @model_validator(mode="after")
    def order_sections(self):
        # 数据库有 CHECK (end_section >= start_section)，模型偶发倒序时在此纠正
        if self.end_section < self.start_section:
            self.start_section, self.end_section = self.end_section, self.start_section
        return self


def serialize_workbook(path: Path) -> str:
    """Render every sheet as coordinate-tagged text rows for the model to read."""
    sheets = read_sheets(path)
    chunks = []
    used = 0
    truncated = False
    for index, sheet in enumerate(sheets, 1):
        lines = [f'=== SHEET {index}: "{sheet.title}" (rows={sheet.max_row}, cols={sheet.max_column}) ===']
        if sheet.merged:
            lines.append("MERGED: " + " ".join(sheet.merged))
        for row_index, cells in sheet.rows:
            if row_index > MAX_ROWS_PER_SHEET:
                break
            cell_texts = []
            for col_index, val in sorted(cells.items()):
                text = str(val or "").strip()
                if not text:
                    continue
                # 换行转义为字面 \n，既保留“教师/周次分行”信息又不破坏 ROW 单行结构
                text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\\n")
                coord = f"{get_column_letter(col_index)}{row_index}"
                cell_texts.append(f'{coord}="{text}"')
            if not cell_texts:
                continue
            lines.append(f"ROW {row_index}: " + " ".join(cell_texts))
        block = "\n".join(lines)
        if used + len(block) > AI_MAX_INPUT_CHARS:
            truncated = True
            break
        chunks.append(block)
        used += len(block) + 1
    if truncated:
        chunks.append("... [TRUNCATED]")
    return "\n".join(chunks)


def _extract_json(text: str) -> str:
    """Defensively unwrap code fences or surrounding prose before json.loads."""
    text = (text or "").strip()
    fence = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL)
    if fence:
        text = fence.group(1).strip()
    start, end = text.find("{"), text.rfind("}")
    return text[start:end + 1] if start != -1 and end > start else text


def _coerce_schedule(payload) -> dict | None:
    """Validate model output course by course; skip bad rows instead of failing all."""
    if not isinstance(payload, dict) or not isinstance(payload.get("courses"), list):
        return None
    courses = []
    seen = set()
    for item in payload["courses"]:
        if not isinstance(item, dict):
            continue
        try:
            course = AICourse.model_validate(item)
        except ValidationError:
            continue
        if not course.name:
            continue
        data = course.model_dump()
        key = (data["name"], data["weekday"], data["start_section"], data["end_section"], tuple(data["weeks"]))
        if key in seen:
            continue
        seen.add(key)
        data["color"] = course_color(data["weekday"], data["start_section"])
        courses.append(data)
    # term 是 (user_id, term) 幂等覆盖键，必须稳定，不能落到随机文件名上
    term = re.sub(r"\s+", " ", str(payload.get("term") or "")).strip()[:80] or "导入课表"
    name = re.sub(r"\s+", " ", str(payload.get("name") or "")).strip()[:80] or "我的课表"
    return {"name": name, "term": term, "courses": merge_section_courses(courses)}


def _resolve_llm_config() -> dict:
    """llm_configs 数据库配置优先，缺失字段回落到环境变量；未配置返回空 dict。"""
    try:
        config = get_llm_config("schedule_import")
    except RuntimeError:
        return {}
    resolved = {
        "base_url": (config.get("base_url") or "").rstrip("/"),
        "api_key": config.get("api_key") or "",
        "model": config.get("model") or "",
        "timeout": float(config.get("timeout_seconds") or AI_TIMEOUT),
        "enable_thinking": bool(config.get("enable_thinking", False)),
    }
    if not (resolved["base_url"] and resolved["api_key"] and resolved["model"]):
        return {}
    return resolved


def parse_with_ai(path: Path) -> tuple[dict | None, str]:
    """Parse an Excel timetable through an OpenAI-compatible chat model.

    Never raises: every failure mode is reported through the engine tag so the
    caller can silently fall back to the deterministic parser.
    模型配置支持管理后台 llm_configs 的 schedule_import 分组动态覆盖。
    """
    config = _resolve_llm_config()
    if not config:
        return None, "ai-not-configured"
    try:
        serialized = serialize_workbook(path)
    except Exception as error:
        logger.warning("AI 输入序列化失败: %s", error.__class__.__name__)
        return None, "ai-unreadable"
    if not serialized.strip():
        return None, "ai-empty-workbook"
    body = {
        "model": config["model"],
        "temperature": 0,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": serialized},
        ],
    }
    if config["enable_thinking"]:
        body["enable_thinking"] = True
    try:
        response = httpx.post(
            f"{config['base_url']}/chat/completions",
            headers={"Authorization": f"Bearer {config['api_key']}"},
            timeout=config["timeout"],
            json=body,
        )
    except httpx.TimeoutException:
        return None, "ai-timeout"
    except Exception as error:
        # 契约是“永不抛出”，因此不只捕 httpx.HTTPError：非法 URL、不支持的协议、
        # 重定向与解码异常等也应降级为兜底解析，而不是变成 500
        logger.warning("AI 服务调用失败: %s", error.__class__.__name__)
        return None, "ai-unavailable"
    if response.status_code != 200:
        return None, f"ai-http-{response.status_code}"
    try:
        content = response.json()["choices"][0]["message"]["content"]
    except (ValueError, KeyError, IndexError, TypeError):
        return None, "ai-invalid-response"
    try:
        payload = json.loads(_extract_json(content))
    except ValueError:
        return None, "ai-invalid-json"
    schedule = _coerce_schedule(payload)
    if not schedule:
        return None, "ai-invalid-schema"
    if not schedule["courses"]:
        return None, "ai-empty"
    return schedule, "ai"


HTML_SYSTEM_PROMPT = """你是高校教务课表网页结构化解析器。用户会提供从教务系统页面抓取的 HTML 表格或文本。

你的任务是从中提取完整的课程安排，并且只输出一个 JSON 对象，禁止输出 markdown 代码块、注释或任何解释文字。

输出 JSON 结构：
{
  "name": "课表名称，如“张三的课表”；无法判断时填“教务课表”",
  "term": "学期标识，如“2025-2026学年第1学期”或“2025-2026-1”；无法判断时填“导入课表”",
  "courses": [
    {
      "name": "课程名称，必填",
      "teacher": "授课教师，缺失填空字符串",
      "room": "教室或上课地点，缺失填空字符串",
      "weekday": 1,
      "start_section": 1,
      "end_section": 2,
      "weeks": "2-16(双)"
    }
  ]
}

解析规则：
1. weekday：1=周一、2=周二 …… 7=周日。
2. start_section / end_section：1 到 12 的整数，取自节次（如“1-2节”取 start=1, end=2）。
3. 连堂课必须合并成一条记录（如 1-2 节与 3-4 节连续且为同一课程教室则合并为 1-4 节）。
4. weeks：原样抄录表中的周次表达式字符串，不要展开成数字数组。例如 "1-16"、"2-16(双)"、"1-8,10-15" 等。缺失时填空字符串。
5. 教室必须保留完整全称（包含教学楼与房间号）；teacher / room 缺失时输出空字符串。
6. 忽略表头、说明、备注、统计、未选课等无关文本。
"""


def clean_html_for_llm(raw_text: str) -> str:
    """去除 script/style/svg/base64 等冗余，精简 HTML/表格文本至合理大小。"""
    raw_text = (raw_text or "").strip()
    if raw_text.startswith("[") or raw_text.startswith("{"):
        try:
            items = json.loads(raw_text)
            if isinstance(items, list):
                chunks = []
                for it in items:
                    title = it.get("title", "")
                    tb = it.get("tablesHtml") or it.get("html", "")
                    txt = it.get("text", "")
                    chunks.append(f"PAGE: {title}\n{tb or txt}")
                raw_text = "\n---\n".join(chunks)
        except Exception:
            pass

    raw_text = re.sub(r"<(script|style|svg|noscript)[^>]*>.*?</\1>", "", raw_text, flags=re.DOTALL | re.IGNORECASE)
    raw_text = re.sub(r'data:[^;]+;base64,[A-Za-z0-9+/=]+', '', raw_text)
    raw_text = re.sub(r'<!--.*?-->', '', raw_text, flags=re.DOTALL)
    raw_text = re.sub(r'[ \t]+', ' ', raw_text)
    raw_text = re.sub(r'\n\s*\n+', '\n', raw_text)
    return raw_text[:AI_MAX_INPUT_CHARS]


def parse_html_with_ai(payload: str) -> tuple[dict | None, str]:
    """通过大模型解析教务系统的 HTML 课表数据（作为确定性解析失败后的兜底）。"""
    config = _resolve_llm_config()
    if not config:
        return None, "ai-not-configured"
    cleaned = clean_html_for_llm(payload)
    if not cleaned.strip():
        return None, "ai-empty-input"

    body = {
        "model": config["model"],
        "temperature": 0,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": HTML_SYSTEM_PROMPT},
            {"role": "user", "content": cleaned},
        ],
    }
    if config["enable_thinking"]:
        body["enable_thinking"] = True

    try:
        response = httpx.post(
            f"{config['base_url']}/chat/completions",
            headers={"Authorization": f"Bearer {config['api_key']}"},
            timeout=config["timeout"],
            json=body,
        )
    except httpx.TimeoutException:
        return None, "ai-timeout"
    except Exception as error:
        logger.warning("HTML AI 服务调用失败: %s", error.__class__.__name__)
        return None, "ai-unavailable"

    if response.status_code != 200:
        return None, f"ai-http-{response.status_code}"

    try:
        content = response.json()["choices"][0]["message"]["content"]
    except (ValueError, KeyError, IndexError, TypeError):
        return None, "ai-invalid-response"

    try:
        payload_data = json.loads(_extract_json(content))
    except ValueError:
        return None, "ai-invalid-json"

    schedule = _coerce_schedule(payload_data)
    if not schedule:
        return None, "ai-invalid-schema"
    if not schedule["courses"]:
        return None, "ai-empty"

    return schedule, "ai-html"
