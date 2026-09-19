"""统一的 Excel 读取层。

屏蔽 .xlsx/.xlsm（OOXML）、.xls（BIFF）以及教务系统常见的"假 .xls"
（实为 HTML 表格、CSV 或改了后缀的 xlsx）之间的差异，向 AI 序列化层与
本地兜底解析器输出同一种规范化结构。

后缀只用于前端与接口的粗筛，真正决定读法的是 detect_format 的 magic bytes。
"""
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from html.parser import HTMLParser
from pathlib import Path

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SUPPORTED_SUFFIXES = {".xls", ".xlsx", ".xlsm"}

ZIP_MAGIC = b"PK\x03\x04"
OLE2_MAGIC = b"\xd0\xcf\x11\xe0"  # BIFF 容器前 4 字节即唯一标识

# 单表行数上限，防止带大量尾部格式的空表拖垮序列化
MAX_ROWS_PER_SHEET = 1500
MAX_OOXML_FILES = 2000
MAX_OOXML_UNCOMPRESSED_BYTES = 64 * 1024 * 1024


def validate_excel_container(path: Path, suffix: str) -> str:
    """校验真实文件类型，并限制 OOXML 解压规模以避免压缩炸弹。

    兼容小程序场景：教务系统导出的"假 .xls"可能是 HTML 表格或 CSV，一并放行。
    """
    kind = detect_format(path)
    if suffix in {".xlsx", ".xlsm"} and kind != "ooxml":
        raise ValueError("文件内容与 Excel 扩展名不匹配")
    if suffix == ".xls" and kind not in {"biff", "ooxml", "html", "csv"}:
        raise ValueError("无法识别为有效的 Excel 文件")
    if kind == "ooxml":
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                total = sum(item.file_size for item in entries)
                if len(entries) > MAX_OOXML_FILES or total > MAX_OOXML_UNCOMPRESSED_BYTES:
                    raise ValueError("Excel 解压后的内容过大")
                for item in entries:
                    if item.compress_size and item.file_size / item.compress_size > 200:
                        raise ValueError("Excel 包含异常压缩内容")
        except zipfile.BadZipFile as error:
            raise ValueError("Excel 文件结构已损坏") from error
    return kind


@dataclass
class Sheet:
    """一张表的规范化视图。

    rows 为 (真实行号, {1-based 列号: 文本})，已跳过全空行；
    merged 为 "A1:H1" 形式的合并区域，HTML/CSV 无此概念时为空列表。
    """
    title: str
    rows: list[tuple[int, dict[int, str]]] = field(default_factory=list)
    merged: list[str] = field(default_factory=list)
    max_row: int = 0
    max_column: int = 0


def text_of(value) -> str:
    """把任意单元格值规范化成文本。

    整数型浮点必须去掉 .0：xlrd 把 BIFF 数字一律读成 float，
    若原样输出 "1.0-2.0"，节次正则 (\\d+)\\s*-\\s*(\\d+) 会错配成 0-2。
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, datetime):
        if not (value.hour or value.minute or value.second):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value).strip()


def detect_format(path: Path) -> str:
    """按 magic bytes 判别真实格式：ooxml / biff / html / csv / empty / unreadable。"""
    try:
        with path.open("rb") as handle:
            head = handle.read(8192)
    except OSError:
        return "unreadable"
    if not head.strip():
        return "empty"
    if head.startswith(ZIP_MAGIC):
        return "ooxml"
    if head.startswith(OLE2_MAGIC):
        return "biff"
    probe = head.decode("utf-8", errors="ignore").lower()
    if "<html" in probe or "<table" in probe or "<!doctype html" in probe:
        return "html"
    return "csv"


def _decode(path: Path) -> str:
    """按常见中文编码依次尝试解码纯文本类"Excel"。"""
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030", "big5", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_ooxml(path: Path, data_only: bool) -> list[Sheet]:
    sheets = []
    workbook = load_workbook(path, data_only=data_only)
    try:
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(max_row=min(worksheet.max_row or 1, MAX_ROWS_PER_SHEET)):
                cells = {}
                for cell in row:
                    text = text_of(cell.value)
                    if text:
                        cells[cell.column] = text
                if cells:
                    rows.append((row[0].row, cells))
            sheets.append(Sheet(
                title=worksheet.title,
                rows=rows,
                merged=[str(item) for item in worksheet.merged_cells.ranges],
                max_row=worksheet.max_row or 0,
                max_column=worksheet.max_column or 0,
            ))
    finally:
        workbook.close()
    return sheets


def _biff_value(worksheet, row: int, column: int, datemode: int):
    cell = worksheet.cell(row, column)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        except xlrd.XLDateError:
            return cell.value
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return None
    return cell.value


def _read_biff(path: Path) -> list[Sheet]:
    try:
        # formatting_info=True 才会解析出 merged_cells
        book = xlrd.open_workbook(str(path), formatting_info=True)
    except Exception:
        # 个别文件的格式记录损坏，合并单元格信息可以缺省，别让整个导入失败
        book = xlrd.open_workbook(str(path))
    sheets = []
    for worksheet in book.sheets():
        rows = []
        for index in range(min(worksheet.nrows, MAX_ROWS_PER_SHEET)):
            cells = {}
            for column in range(worksheet.ncols):
                text = text_of(_biff_value(worksheet, index, column, book.datemode))
                if text:
                    cells[column + 1] = text
            if cells:
                rows.append((index + 1, cells))
        # xlrd 的 merged_cells 是 (rlo, rhi, clo, chi) 且上界开区间
        merged = [
            f"{get_column_letter(clo + 1)}{rlo + 1}:{get_column_letter(chi)}{rhi}"
            for rlo, rhi, clo, chi in worksheet.merged_cells
        ]
        sheets.append(Sheet(worksheet.name, rows, merged, worksheet.nrows, worksheet.ncols))
    return sheets


def _span_of(attributes: dict, key: str) -> int:
    try:
        return max(1, int(str(attributes.get(key) or 1).strip()))
    except ValueError:
        return 1


class _TableParser(HTMLParser):
    """抓取最外层 <table> 的单元格文本与 colspan/rowspan，嵌套表忽略。"""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.grids = []
        self.depth = 0
        self.grid = None
        self.row = None
        self.text = None
        self.span = (1, 1)

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.depth += 1
            if self.depth == 1:
                self.grid = []
            return
        if self.depth != 1:
            return
        if tag == "tr":
            self.row = []
        elif tag in ("td", "th"):
            attributes = dict(attrs)
            self.span = (_span_of(attributes, "colspan"), _span_of(attributes, "rowspan"))
            self.text = []
        elif tag == "br" and self.text is not None:
            self.text.append("\n")

    def handle_endtag(self, tag):
        if tag == "table":
            if self.depth == 1 and self.grid is not None:
                self.grids.append(self.grid)
                self.grid = None
            self.depth = max(0, self.depth - 1)
            return
        if self.depth != 1:
            return
        if tag in ("td", "th") and self.text is not None:
            if self.row is not None:
                self.row.append(("".join(self.text).strip(), *self.span))
            self.text = None
        elif tag == "tr" and self.row is not None and self.grid is not None:
            self.grid.append(self.row)
            self.row = None

    def handle_data(self, data):
        if self.text is not None:
            self.text.append(data)


def _grid_to_rows(grid: list[list[tuple[str, int, int]]]) -> tuple[list[tuple[int, dict[int, str]]], int]:
    """把带 colspan/rowspan 的单元格铺进二维矩阵，返回规范化 rows 与最大列号。"""
    occupied: set[tuple[int, int]] = set()
    matrix: dict[int, dict[int, str]] = {}
    for row_index, row in enumerate(grid):
        column = 0
        for text, colspan, rowspan in row:
            while (row_index, column) in occupied:
                column += 1
            for down in range(rowspan):
                for across in range(colspan):
                    occupied.add((row_index + down, column + across))
            if text:
                matrix.setdefault(row_index + 1, {})[column + 1] = text
            column += colspan
    rows = [(row_index, cells) for row_index, cells in sorted(matrix.items()) if cells]
    max_column = max((max(cells) for _, cells in rows), default=0)
    return rows[:MAX_ROWS_PER_SHEET], max_column


def _read_html(path: Path) -> list[Sheet]:
    parser = _TableParser()
    parser.feed(_decode(path))
    parser.close()
    sheets = []
    for index, grid in enumerate(parser.grids, 1):
        rows, max_column = _grid_to_rows(grid)
        if not rows:
            continue
        sheets.append(Sheet(f"表格{index}", rows, [], rows[-1][0], max_column))
    if not sheets:
        # 没有 <table> 的 HTML：退化成按行文本，交给 AI 自己判断
        lines = [line.strip() for line in _decode(path).splitlines()]
        rows = [(index, {1: re.sub(r"<[^>]+>", "", line)})
                for index, line in enumerate(lines, 1) if re.sub(r"<[^>]+>", "", line).strip()]
        if rows:
            sheets.append(Sheet(path.stem, rows[:MAX_ROWS_PER_SHEET], [], len(rows), 1))
    return sheets


def _read_csv(path: Path) -> list[Sheet]:
    text = _decode(path)
    head = text[:8192]
    try:
        delimiter = csv.Sniffer().sniff(head, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = "\t" if head.count("\t") > head.count(",") else ","
    rows = []
    for index, record in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), 1):
        cells = {column: value.strip() for column, value in enumerate(record, 1) if value and value.strip()}
        if cells:
            rows.append((index, cells))
        if index >= MAX_ROWS_PER_SHEET:
            break
    max_column = max((max(cells) for _, cells in rows), default=0)
    return [Sheet(path.stem, rows, [], len(rows), max_column)] if rows else []


def content_chars(sheets: list[Sheet]) -> int:
    """统计规范化结果里的有效字符数，用于判断是否读到了实质内容。"""
    return sum(len(text) for sheet in sheets for _, cells in sheet.rows for text in cells.values())


def read_sheets(path: Path, kind: str | None = None) -> list[Sheet]:
    """按真实格式读取工作簿，返回统一的 Sheet 列表。"""
    kind = kind or detect_format(path)
    if kind == "ooxml":
        sheets = _read_ooxml(path, data_only=True)
        if content_chars(sheets) < 50:
            # 脚本生成、未经 Excel 保存的工作簿没有公式缓存值，退回读取公式本身
            sheets = _read_ooxml(path, data_only=False)
        return sheets
    if kind == "biff":
        return _read_biff(path)
    if kind == "html":
        return _read_html(path)
    if kind == "csv":
        return _read_csv(path)
    return []
