"""Regression tests for deterministic schedule parsing and normalization."""
from openpyxl import Workbook

from app.parser import normalize_courses, parse_excel_schedule, parse_grid_schedule


def test_normalize_courses_rules_and_deduplication():
    raw_courses = [
        # 正常课程
        {"name": " 高等数学 ", "teacher": " 张三 ", "room": " A101 ", "weekday": 1,
         "start_section": 1, "end_section": 2, "weeks": [1, 2, 30, 31]},
        # 相同去重键 (name, weekday, start, end, weeks)，应去重
        {"name": "高等数学", "teacher": "其他老师", "room": "A102", "weekday": 1,
         "start_section": 1, "end_section": 2, "weeks": [1, 2, 30]},
        # 连堂课相邻节次，应与高等数学合并为 1-4
        {"name": "高等数学", "teacher": " 张三 ", "room": " A101 ", "weekday": 1,
         "start_section": 3, "end_section": 4, "weeks": [1, 2, 30]},
        # 节次反序，自动调整
        {"name": "大学英语", "weekday": 2, "start_section": 4, "end_section": 3, "weeks": "1-8周"},
        # 无效名称，应丢弃
        {"name": "   ", "weekday": 3, "start_section": 1, "end_section": 2},
        # 星期超出 1~7，应丢弃
        {"name": "无效星期", "weekday": 8, "start_section": 1, "end_section": 2},
        # 节次超出 1~12，应丢弃
        {"name": "无效节次", "weekday": 3, "start_section": 0, "end_section": 2},
    ]
    normalized = normalize_courses(raw_courses)
    assert len(normalized) == 2
    math = next(c for c in normalized if c["name"] == "高等数学")
    assert math["weekday"] == 1
    assert math["start_section"] == 1
    assert math["end_section"] == 4
    assert math["weeks"] == [1, 2, 30]  # 31 被剔除
    assert math["teacher"] == "张三"
    assert math["room"] == "A101"

    english = next(c for c in normalized if c["name"] == "大学英语")
    assert english["start_section"] == 3
    assert english["end_section"] == 4
    assert english["weeks"] == list(range(1, 9))


def test_parse_grid_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "2026-2027学年第1学期"
    for col, title in enumerate(["节次", "星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"], 1):
        ws.cell(row=2, column=col, value=title)
    ws["A3"] = "第1节"
    ws["B3"] = "物理\n钱七\nB101\n1-16周"
    ws["C3"] = "化学\n孙八\nB102\n1-8周"
    ws["A4"] = "第2节"
    ws["B4"] = "物理\n钱七\nB101\n1-16周"
    target = tmp_path / "grid.xlsx"
    wb.save(target)
    parsed = parse_grid_schedule(target)
    assert parsed and len(parsed["courses"]) == 2
    physics = next(c for c in parsed["courses"] if c["name"] == "物理")
    assert (physics["start_section"], physics["end_section"]) == (1, 2)
    assert physics["weeks"] == list(range(1, 17))

    # parse_excel_schedule 应当自动回退到 grid
    auto_parsed = parse_excel_schedule(target)
    assert auto_parsed and len(auto_parsed["courses"]) == 2


def test_parse_grid_numeric_sections_and_multi_courses(tmp_path):
    """测试强智等系统导出的数字节次(1.0, 2.0)与单单元格多课程(含换行后教室)正确解析与合并。"""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "2026-2027学年第1学期 谭锃 课表"
    for col, title in enumerate(["节次", "星期一", "星期二", "星期三", "星期四", "星期五"], 1):
        ws.cell(row=2, column=col, value=title)
    # 节次为浮点数 1.0, 2.0
    ws["A3"] = 1.0
    ws["B3"] = "思想概论\n\n杨老师【1-8周】\n7-北201\n\n思想概论\n\n杨老师【9-16周】\nVR教室\n"
    ws["A4"] = 2.0
    ws["B4"] = "思想概论\n\n杨老师【1-8周】\n7-北201\n\n思想概论\n\n杨老师【9-16周】\nVR教室\n"

    target = tmp_path / "xskb_sample.xlsx"
    wb.save(target)
    parsed = parse_grid_schedule(target)
    assert parsed is not None
    assert parsed["term"] == "2026-2027学年第1学期"
    courses = parsed["courses"]
    assert len(courses) == 2

    c1 = next(c for c in courses if c["room"] == "7-北201")
    assert c1["name"] == "思想概论"
    assert c1["teacher"] == "杨老师"
    assert c1["start_section"] == 1
    assert c1["end_section"] == 2
    assert c1["weeks"] == list(range(1, 9))

    c2 = next(c for c in courses if c["room"] == "VR教室")
    assert c2["name"] == "思想概论"
    assert c2["teacher"] == "杨老师"
    assert c2["start_section"] == 1
    assert c2["end_section"] == 2
    assert c2["weeks"] == list(range(9, 17))
