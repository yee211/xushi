"""Unit tests for schedule import and overwrite cascade deletion."""
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.routers.importer as importer_module
from app.auth import get_current_user
from app.main import app


class Result:
    def __init__(self, row=None):
        self.row = row

    def fetchone(self):
        return self.row


class MockDb:
    def __init__(self):
        self.deleted_adjusted = None
        self.deleted_courses = None
        self.existing_schedule = {"id": 11, "name": "旧课表", "term": "2026-2027-1学期"}

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def execute(self, sql, params=None):
        if "pg_advisory_xact_lock" in sql:
            return Result()
        if "SELECT * FROM schedules" in sql:
            return Result(self.existing_schedule)
        if "DELETE FROM schedules" in sql:
            self.deleted_adjusted = params
            return Result()
        if "UPDATE schedules" in sql:
            return Result({**self.existing_schedule, "name": params[0]})
        if "DELETE FROM courses" in sql:
            self.deleted_courses = params
            return Result()
        raise AssertionError(f"Unexpected SQL: {sql}")

    def cursor(self):
        class MockCursor:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def executemany(self, sql, rows):
                pass
        return MockCursor()


def make_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "2026-2027学年第1学期课程明细"
    headers = ["星期", "节次", "课程名称", "教师", "周次", "教室"]
    for col, title in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=title)
    ws.cell(row=3, column=1, value="星期一")
    ws.cell(row=3, column=2, value="1-2")
    ws.cell(row=3, column=3, value="高等数学")
    ws.cell(row=3, column=4, value="张三")
    ws.cell(row=3, column=5, value="1-16周")
    ws.cell(row=3, column=6, value="A101")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_overwrite_deletes_adjusted_variant(monkeypatch):
    db = MockDb()
    monkeypatch.setattr(importer_module, "connect", lambda: db)
    monkeypatch.setattr(importer_module, "parse_with_ai", lambda path: (None, "ai-not-configured"))
    app.dependency_overrides[get_current_user] = lambda: {"id": 7, "username": "u7", "email": "u7@example.com"}

    client = TestClient(app)
    try:
        xlsx = make_xlsx_bytes()
        res = client.post(
            "/api/import",
            files={"file": ("schedule.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert res.status_code == 200
        assert db.deleted_adjusted == (11, 7)
        assert db.deleted_courses == (11,)
    finally:
        app.dependency_overrides.pop(get_current_user, None)
