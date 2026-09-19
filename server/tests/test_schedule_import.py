"""Regression tests for schedule import and overwrite pipeline."""
import io
from contextlib import nullcontext
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.auth import get_current_user as current_user
from app.main import app
from app.parser import normalize_courses
from app.routers import importer as main
from app.services.schedule_import import write_schedule


class Result:
    def __init__(self, row=None):
        self.row = row

    def fetchone(self):
        return self.row


class Cursor:
    def __init__(self, db):
        self.db = db

    def executemany(self, sql, rows):
        self.db.inserted = list(rows)


class ImportDb:
    """原有 write_schedule 单元测试所用的 mock DB。"""
    def __init__(self):
        self.schedule = {"id": 11, "name": "旧课表", "term": "2026秋"}
        self.deleted_adjusted = None
        self.deleted_courses = None
        self.inserted = []

    def execute(self, sql, params=None):
        if "pg_advisory_xact_lock" in sql:
            return Result()
        if "SELECT * FROM schedules" in sql:
            return Result(self.schedule)
        if "DELETE FROM schedules" in sql:
            self.deleted_adjusted = params
            return Result()
        if "UPDATE schedules" in sql:
            return Result({**self.schedule, "name": params[0], "term": params[1]})
        if "DELETE FROM courses" in sql:
            self.deleted_courses = params
            return Result()
        raise AssertionError(sql)

    def cursor(self):
        return nullcontext(Cursor(self))


class MockImportDb:
    """支持事务生命周期、级联清理及失败回滚的 Mock 数据库。"""
    def __init__(self, existing_schedule=None, fail_courses_insert=False):
        self.existing_schedule = existing_schedule
        self.fail_courses_insert = fail_courses_insert
        self.deleted_adjusted = None
        self.deleted_courses = None
        self.inserted_schedule = None
        self.updated_schedule = None
        self.inserted_courses = []
        self.rolled_back = False
        self.committed = False
        self.next_id = 100

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        if exc_type is not None:
            self.rolled_back = True
        else:
            self.committed = True
        return False

    def execute(self, sql, params=None):
        if "pg_advisory_xact_lock" in sql:
            return Result()
        if "SELECT * FROM schedules" in sql:
            return Result(self.existing_schedule)
        if "DELETE FROM schedules" in sql:
            self.deleted_adjusted = params
            return Result()
        if "UPDATE schedules" in sql:
            self.updated_schedule = params
            return Result({
                "id": self.existing_schedule["id"],
                "name": params[0],
                "term": params[1],
                "start_date": params[2],
                "end_date": params[3],
                "variant_type": "original",
            })
        if "INSERT INTO schedules" in sql:
            self.inserted_schedule = params
            self.next_id += 1
            return Result({
                "id": self.next_id,
                "name": params[1],
                "term": params[2],
                "start_date": params[3],
                "end_date": params[4],
                "variant_type": "original",
            })
        if "DELETE FROM courses" in sql:
            self.deleted_courses = params
            return Result()
        raise AssertionError(f"Unexpected SQL: {sql}")

    def cursor(self):
        class MockCursor:
            def __init__(self, db):
                self.db = db

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def executemany(self, sql, rows):
                if self.db.fail_courses_insert:
                    raise RuntimeError("DB insert courses failed")
                self.db.inserted_courses.extend(rows)

        return MockCursor(self)


def make_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "2026-2027学年第1学期课程明细"
    headers = ["星期", "节次", "课程名称", "教师", "周次", "教室"]
    for col, title in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=title)
    ws.cell(row=3, column=1, value="星期一")
    ws.cell(row=3, column=2, value="1-2")
    ws.cell(row=3, column=3, value="高等数学")
    ws.cell(row=3, column=4, value="张三")
    ws.cell(row=3, column=5, value="1-16周")
    ws.cell(row=3, column=6, value="A101")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HTML_XLS_BYTES = (
    "<html><body><table>"
    "<tr><td>星期</td><td>节次</td><td>课程</td><td>教师</td><td>周次</td><td>教室</td></tr>"
    "<tr><td>星期一</td><td>1-2</td><td>大学物理</td><td>李四</td><td>1-16周</td><td>B202</td></tr>"
    "</table></body></html>"
).encode()

CSV_XLS_BYTES = "星期,节次,课程,教师,周次,教室\n星期二,3-4,线性代数,王五,1-16周,C303\n".encode("gb18030")


@pytest.fixture(autouse=True)
def disable_real_ai_network(monkeypatch):
    """单元测试默认不发起真实 AI 外部网络请求，默认返回 ai-not-configured 快速走本地解析。"""
    monkeypatch.setattr(main, "parse_with_ai", lambda path: (None, "ai-not-configured"))


@pytest.fixture
def client():
    app.dependency_overrides[current_user] = lambda: {"id": 7, "openid": "test_user"}
    yield TestClient(app)
    app.dependency_overrides.pop(current_user, None)


def test_overwrite_deletes_adjusted_copy_before_replacing_courses():
    db = ImportDb()
    result = write_schedule(db, user_id=7, overwrite=True, parsed={
        "name": "新课表",
        "term": "2026秋",
        "courses": [{"name": "高等数学", "weekday": 1, "start_section": 1,
                     "end_section": 2, "weeks": [1, 2]}],
    })

    assert db.deleted_adjusted == (11, 7)
    assert db.deleted_courses == (11,)
    assert result["imported"] == 1 and result["schedule_id"] == 11 and result["replaced"] is True


def test_normalize_courses_rules_and_deduplication():
    raw_courses = [
        # 正常课程
        {"name": " 高等数学 ", "teacher": " 张三 ", "room": " A101 ", "weekday": 1,
         "start_section": 1, "end_section": 2, "weeks": [1, 2, 30, 31]},
        # 相同去重键 (name, weekday, start, end, weeks)，应去重
        {"name": "高等数学", "teacher": "其他老师", "room": "A102", "weekday": 1,
         "start_section": 1, "end_section": 2, "weeks": [1, 2, 30]},
        # 连堂课相邻节次，应与高等数学合并为 1-4
        {"name": "高等数学", "teacher": " 张三 ", "room": " A101 ", "weekday": 1,
         "start_section": 3, "end_section": 4, "weeks": [1, 2, 30]},
        # 节次反序，自动调整
        {"name": "大学英语", "weekday": 2, "start_section": 4, "end_section": 3, "weeks": "1-8周"},
        # 无效名称，应丢弃
        {"name": "   ", "weekday": 3, "start_section": 1, "end_section": 2},
        # 星期超出 1~7，应丢弃
        {"name": "无效星期", "weekday": 8, "start_section": 1, "end_section": 2},
        # 节次超出 1~12，应丢弃
        {"name": "无效节次", "weekday": 3, "start_section": 0, "end_section": 2},
    ]
    normalized = normalize_courses(raw_courses)
    assert len(normalized) == 2
    math = next(c for c in normalized if c["name"] == "高等数学")
    assert math["weekday"] == 1
    assert math["start_section"] == 1
    assert math["end_section"] == 4
    assert math["weeks"] == [1, 2, 30]  # 31 被剔除
    assert math["teacher"] == "张三"
    assert math["room"] == "A101"

    english = next(c for c in normalized if c["name"] == "大学英语")
    assert english["start_section"] == 3
    assert english["end_section"] == 4
    assert english["weeks"] == list(range(1, 9))


def test_write_schedule_empty_courses_raises_422():
    db = ImportDb()
    with pytest.raises(HTTPException) as exc:
        write_schedule(db, user_id=7, parsed={"name": "空课表", "courses": []})
    assert exc.value.status_code == 422


def test_import_new_schedule_success(client, monkeypatch):
    db = MockImportDb(existing_schedule=None)
    monkeypatch.setattr(main, "connect", lambda: db)
    xlsx_bytes = make_xlsx_bytes()

    response = client.post(
        "/api/import",
        files={"file": ("schedule.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"term_name": "2026-2027学年第1学期", "overwrite": "false"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["imported"] == 1
    assert data["schedule_id"] == 101
    assert data["replaced"] is False
    assert data["suggestions"] == []
    assert db.inserted_schedule is not None
    assert db.inserted_schedule[1] == "2026-2027学年第1学期"
    assert len(db.inserted_courses) == 1


def test_import_duplicate_without_overwrite_returns_409(client, monkeypatch):
    existing = {"id": 11, "name": "2026-2027学年第1学期", "term": "2026-2027学年第1学期"}
    db = MockImportDb(existing_schedule=existing)
    monkeypatch.setattr(main, "connect", lambda: db)
    xlsx_bytes = make_xlsx_bytes()

    response = client.post(
        "/api/import",
        files={"file": ("schedule.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"term_name": "2026-2027学年第1学期", "overwrite": "false"},
    )
    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["code"] == "schedule_exists"
    assert detail["schedule_id"] == 11
    assert db.deleted_courses is None
    assert db.deleted_adjusted is None


def test_import_overwrite_keeps_id_and_deletes_adjusted_copy(client, monkeypatch):
    existing = {"id": 11, "name": "2026-2027学年第1学期", "term": "2026-2027学年第1学期"}
    db = MockImportDb(existing_schedule=existing)
    monkeypatch.setattr(main, "connect", lambda: db)
    xlsx_bytes = make_xlsx_bytes()

    response = client.post(
        "/api/import",
        files={"file": ("schedule.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"term_name": "2026-2027学年第1学期", "overwrite": "true"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["schedule_id"] == 11
    assert data["replaced"] is True
    # 验证旧调课版课表被删除
    assert db.deleted_adjusted == (11, 7)
    # 验证旧课程被删除并重新插入
    assert db.deleted_courses == (11,)
    assert len(db.inserted_courses) == 1


def test_import_xlsx_ai_success(client, monkeypatch):
    db = MockImportDb(existing_schedule=None)
    monkeypatch.setattr(main, "connect", lambda: db)
    xlsx_bytes = make_xlsx_bytes()

    monkeypatch.setattr(main, "parse_with_ai", lambda path: ({
        "name": "AI解析课表",
        "term": "2026秋",
        "courses": [{"name": "高等数学", "weekday": 1, "start_section": 1, "end_section": 2, "weeks": [1, 2]}],
    }, "ai"))

    response = client.post(
        "/api/import",
        files={"file": ("schedule.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"overwrite": "false"},
    )
    assert response.status_code == 200
    assert response.json()["engine"] == "ai"


def test_import_xlsx_ai_failure_local_fallback(client, monkeypatch):
    db = MockImportDb(existing_schedule=None)
    monkeypatch.setattr(main, "connect", lambda: db)
    xlsx_bytes = make_xlsx_bytes()

    monkeypatch.setattr(main, "parse_with_ai", lambda path: (None, "ai-timeout"))

    response = client.post(
        "/api/import",
        files={"file": ("schedule.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"overwrite": "false"},
    )
    assert response.status_code == 200
    assert response.json()["engine"] == "excel-fallback(ai-timeout)"


def test_import_xls_attempts_ai(client, monkeypatch):
    db = MockImportDb(existing_schedule=None)
    monkeypatch.setattr(main, "connect", lambda: db)

    ai_called_path = None

    def fake_parse_ai(path):
        nonlocal ai_called_path
        ai_called_path = path
        return None, "ai-not-configured"

    monkeypatch.setattr(main, "parse_with_ai", fake_parse_ai)

    response = client.post(
        "/api/import",
        files={"file": ("export.xls", HTML_XLS_BYTES, "application/vnd.ms-excel")},
        data={"overwrite": "false"},
    )
    assert response.status_code == 200
    assert ai_called_path is not None
    assert str(ai_called_path).endswith(".xls")
    assert response.json()["engine"] == "excel-fallback(ai-not-configured)"


def test_import_fake_xls_html_and_csv(client, monkeypatch):
    monkeypatch.setattr(main, "parse_with_ai", lambda path: (None, "ai-not-configured"))

    # HTML 伪 .xls
    db_html = MockImportDb(existing_schedule=None)
    monkeypatch.setattr(main, "connect", lambda: db_html)
    resp_html = client.post(
        "/api/import",
        files={"file": ("export_html.xls", HTML_XLS_BYTES, "application/vnd.ms-excel")},
        data={"overwrite": "false"},
    )
    assert resp_html.status_code == 200
    assert resp_html.json()["imported"] == 1
    assert db_html.inserted_courses[0][1] == "大学物理"

    # CSV 伪 .xls
    db_csv = MockImportDb(existing_schedule=None)
    monkeypatch.setattr(main, "connect", lambda: db_csv)
    resp_csv = client.post(
        "/api/import",
        files={"file": ("export_csv.xls", CSV_XLS_BYTES, "application/vnd.ms-excel")},
        data={"overwrite": "false"},
    )
    assert resp_csv.status_code == 200
    assert resp_csv.json()["imported"] == 1
    assert db_csv.inserted_courses[0][1] == "线性代数"


def test_import_extension_mismatch_returns_400(client):
    response = client.post(
        "/api/import",
        files={"file": ("fake.xlsx", b"This is plain text, not a zip file", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"overwrite": "false"},
    )
    assert response.status_code == 400
    assert "不匹配" in response.json()["detail"]


def test_import_oversized_file_returns_413(client, monkeypatch):
    monkeypatch.setattr(main, "settings", SimpleNamespace(max_upload_bytes=100))
    response = client.post(
        "/api/import",
        files={"file": ("large.xlsx", b"x" * 200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"overwrite": "false"},
    )
    assert response.status_code == 413


def test_import_no_valid_courses_returns_422(client, monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "只有表头没有课程"
    buf = io.BytesIO()
    wb.save(buf)
    empty_xlsx = buf.getvalue()

    monkeypatch.setattr(main, "parse_with_ai", lambda path: (None, "ai-not-configured"))

    response = client.post(
        "/api/import",
        files={"file": ("empty.xlsx", empty_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"overwrite": "false"},
    )
    assert response.status_code == 422


def test_import_db_failure_rolls_back_transaction(client, monkeypatch):
    db = MockImportDb(existing_schedule=None, fail_courses_insert=True)
    monkeypatch.setattr(main, "connect", lambda: db)
    xlsx_bytes = make_xlsx_bytes()

    with pytest.raises(RuntimeError, match="DB insert courses failed"):
        client.post(
            "/api/import",
            files={"file": ("schedule.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"overwrite": "false"},
        )
    assert db.rolled_back is True
    assert db.committed is False


def test_public_app_config_endpoint(client, monkeypatch):
    monkeypatch.delenv("SHOW_AGENT_ENTRY", raising=False)
    res = client.get("/api/config")
    assert res.status_code == 200
    assert res.json() == {"show_agent": True}

    monkeypatch.setenv("SHOW_AGENT_ENTRY", "false")
    res = client.get("/api/config")
    assert res.status_code == 200
    assert res.json() == {"show_agent": False}


def test_create_demo_schedule_endpoint(client, monkeypatch):
    from app.routers import schedules as schedules_module

    db = MockImportDb(existing_schedule=None)
    monkeypatch.setattr(schedules_module, "connect", lambda: db)
    res = client.post("/api/schedules/demo")
    assert res.status_code == 201
    data = res.json()
    assert data["schedule_id"] == 101
    assert "示例" in data["term_name"]
    assert data["imported"] == 7
    assert len(db.inserted_courses) == 7

